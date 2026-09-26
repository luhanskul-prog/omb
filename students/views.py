from django.shortcuts import render, redirect, get_object_or_404
import csv
import io
from datetime import datetime

from django.contrib import messages
from django.db import transaction
from django.http import HttpResponse
from django.contrib.auth.models import User

from .forms import StudentForm
from .models import Student, OnlineApplication
from accounts.models import UserProfile
from accounts.staff_access import role_permission_required
from parents.models import Parent


def create_parent_account(student):
    """
    Automatically create or update a parent account
    when a learner is registered.
    """

    parent_name = student.parent_name.strip()
    admission_no = student.admission_no.strip()
    parent_phone = student.parent_phone.strip()

    if not parent_name or not admission_no:
        return

    # Username based on the parent name
    username = parent_name.replace(" ", "_")

    # Check whether this parent already has an account
    user = User.objects.filter(username=username).first()

    if not user:
        user = User.objects.create_user(
            username=username,
            password=admission_no
        )

    # Create UserProfile if it does not exist
    profile, created = UserProfile.objects.get_or_create(
        user=user,
        defaults={
            "role": "PARENT",
            "phone": parent_phone,
        }
    )

    if not created:
        profile.role = "PARENT"
        profile.phone = parent_phone
        profile.save()

    # Create Parent profile if it does not exist
    parent, created = Parent.objects.get_or_create(
        user=user,
        defaults={
            "full_name": parent_name,
            "phone": parent_phone,
        }
    )

    if not created:
        parent.full_name = parent_name
        parent.phone = parent_phone
        parent.save()

    # Link the learner to the parent
    parent.children.add(student)



# ============================================================
# STUDENT MANAGEMENT DASHBOARD
# ============================================================


@role_permission_required("view_student", "students")
def student_management(request):

    from django.db.models import Q
    from fees.models import AcademicYear

    students = Student.objects.filter(is_active=True).order_by("admission_no")

    search = request.GET.get("search", "").strip()
    selected_class = request.GET.get("class", "").strip()
    selected_stream = request.GET.get("stream", "").strip()

    if search:
        students = students.filter(
            Q(first_name__icontains=search)
            | Q(middle_name__icontains=search)
            | Q(last_name__icontains=search)
            | Q(admission_no__icontains=search)
        )

    if selected_class:
        students = students.filter(class_name=selected_class)

    if selected_stream:
        students = students.filter(stream=selected_stream)

    from timetable.models import TimetableClass
    from academic.models import Stream

    configured_classes = (
        TimetableClass.objects
        .filter(active=True)
        .exclude(name="")
        .values_list("name", flat=True)
        .distinct()
        .order_by("name")
    )

    classes = [
        (name, name)
        for name in configured_classes
    ]

    streams = (
        Stream.objects
        .filter(is_active=True)
        .exclude(name="")
        .values_list("name", flat=True)
        .distinct()
        .order_by("name")
    )

    gender_choices = Student.GENDER_CHOICES

    academic_years = AcademicYear.objects.filter(
        is_active=True
    ).order_by("-year")

    total_students = Student.objects.filter(is_active=True).count()

    male_students = Student.objects.filter(
        is_active=True,
        gender="Male"
    ).count()

    female_students = Student.objects.filter(
        is_active=True,
        gender="Female"
    ).count()

    class_counts = []

    for code, label in classes:
        count = Student.objects.filter(
            is_active=True,
            class_name=code
        ).count()

        class_counts.append({
            "code": code,
            "label": label,
            "count": count,
        })

    return render(
        request,
        "students/student_management.html",
        {
            "students": students,
            "classes": classes,
            "streams": streams,
            "gender_choices": gender_choices,
            "academic_years": academic_years,
            "search": search,
            "selected_class": selected_class,
            "selected_stream": selected_stream,
            "total_students": total_students,
            "male_students": male_students,
            "female_students": female_students,
            "class_counts": class_counts,
        }
    )

def _clean_import_value(value):
    if value is None:
        return ""

    return str(value).strip()


def _parse_import_date(value):
    value = _clean_import_value(value)

    if not value:
        return None

    formats = [
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%m/%d/%Y",
    ]

    for fmt in formats:
        try:
            return datetime.strptime(
                value,
                fmt
            ).date()
        except ValueError:
            continue

    return None


def _read_student_import_file(uploaded_file):

    filename = uploaded_file.name.lower()

    if filename.endswith(".csv"):

        content = uploaded_file.read().decode(
            "utf-8-sig"
        )

        reader = csv.DictReader(
            io.StringIO(content)
        )

        return list(reader), list(
            reader.fieldnames or []
        )

    if filename.endswith(".xlsx"):

        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ValueError(
                "Excel support requires openpyxl. "
                "Run: python -m pip install openpyxl"
            )

        workbook = load_workbook(
            uploaded_file,
            read_only=True,
            data_only=True
        )

        sheet = workbook.active

        rows = list(
            sheet.iter_rows(
                values_only=True
            )
        )

        if not rows:
            return [], []

        headers = [
            _clean_import_value(value).lower()
            for value in rows[0]
        ]

        records = []

        for row in rows[1:]:

            record = {}

            for index, header in enumerate(headers):

                if header:

                    value = (
                        row[index]
                        if index < len(row)
                        else ""
                    )

                    record[header] = value

            records.append(record)

        return records, headers

    raise ValueError(
        "Unsupported file type. "
        "Please upload CSV or Excel (.xlsx)."
    )


@role_permission_required("view_student", "students")

# ============================================================
# EXPORT STUDENTS
# ============================================================

@role_permission_required("view_student", "students")


def export_students_page(request):
    from fees.models import AcademicYear

    classes = [
        ("PG", "PG"),
        ("PP1", "PP1"),
        ("PP2", "PP2"),
        ("G1", "Grade 1"),
        ("G2", "Grade 2"),
        ("G3", "Grade 3"),
        ("G4", "Grade 4"),
        ("G5", "Grade 5"),
        ("G6", "Grade 6"),
        ("G7", "Grade 7"),
        ("G8", "Grade 8"),
        ("G9", "Grade 9"),
        ("G10", "Grade 10"),
        ("G11", "Grade 11"),
        ("G12", "Grade 12"),
    ]

    streams = (
        Student.objects
        .exclude(stream="")
        .values_list("stream", flat=True)
        .distinct()
        .order_by("stream")
    )

    academic_years = AcademicYear.objects.filter(
        is_active=True
    ).order_by("-year")

    gender_choices = Student.GENDER_CHOICES

    return render(
        request,
        "students/export_students.html",
        {
            "classes": classes,
            "streams": streams,
            "academic_years": academic_years,
            "gender_choices": gender_choices,
        }
    )


def export_students(request):

    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
    except ImportError:
        return HttpResponse(
            "Excel export requires openpyxl. Run: python -m pip install openpyxl"
        )

    from fees.models import AcademicYear
    from django.db.models import Q

    export_type = (request.GET.get("export_type") or "all").strip()
    selected_class = (request.GET.get("class") or "").strip()
    selected_stream = (request.GET.get("stream") or "").strip()
    selected_gender = (request.GET.get("gender") or "").strip()
    selected_year = (request.GET.get("academic_year") or "").strip()

    students = Student.objects.filter(is_active=True)

    # -----------------------------------------------------
    # ACADEMIC YEAR
    # Academic year is linked through FeeRecord.
    # -----------------------------------------------------
    if selected_year:
        students = students.filter(
            fee_records__academic_year_id=selected_year
        ).distinct()

    # -----------------------------------------------------
    # EXPORT BY CLASS
    # -----------------------------------------------------
    if export_type in [
        "class",
        "class_stream",
        "class_gender",
        "class_stream_gender",
    ]:
        if selected_class:
            students = students.filter(
                class_name=selected_class
            )

    # -----------------------------------------------------
    # EXPORT BY STREAM
    # -----------------------------------------------------
    if export_type in [
        "class_stream",
        "class_stream_gender",
    ]:
        if selected_stream:
            students = students.filter(
                stream=selected_stream
            )

    # -----------------------------------------------------
    # EXPORT BY GENDER
    # -----------------------------------------------------
    if export_type in [
        "class_gender",
        "class_stream_gender",
    ]:
        if selected_gender:
            students = students.filter(
                gender=selected_gender
            )

    students = students.order_by(
        "class_name",
        "stream",
        "admission_no"
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Students"

    headers = [
        "No.",
        "Admission Number",
        "Learner Name",
        "Gender",
        "Date of Birth",
        "Class",
        "Stream",
        "Parent Name",
        "Parent Phone",
        "Date Admitted",
    ]

    ws.append(headers)

    for number, student in enumerate(students, start=1):

        ws.append([
            number,
            student.admission_no,
            student.get_full_name(),
            student.gender,
            student.date_of_birth.strftime("%Y-%m-%d")
            if student.date_of_birth else "",
            student.class_name,
            student.stream,
            student.parent_name,
            student.parent_phone,
            student.date_admitted.strftime("%Y-%m-%d")
            if student.date_admitted else "",
        ])

    # Header formatting
    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = {
        1: 8,
        2: 20,
        3: 30,
        4: 12,
        5: 16,
        6: 15,
        7: 15,
        8: 25,
        9: 20,
        10: 16,
    }

    for column, width in widths.items():
        ws.column_dimensions[
            get_column_letter(column)
        ].width = width

    # -----------------------------------------------------
    # FILE NAME
    # -----------------------------------------------------
    filename_parts = ["students"]

    if export_type == "all":
        filename_parts.append("all")

    elif export_type == "class":
        filename_parts.append("by_class")
        if selected_class:
            filename_parts.append(selected_class)

    elif export_type == "class_stream":
        filename_parts.append("class_stream")
        if selected_class:
            filename_parts.append(selected_class)
        if selected_stream:
            filename_parts.append(selected_stream)

    elif export_type == "class_gender":
        filename_parts.append("class_gender")
        if selected_class:
            filename_parts.append(selected_class)
        if selected_gender:
            filename_parts.append(selected_gender)

    elif export_type == "class_stream_gender":
        filename_parts.append("class_stream_gender")
        if selected_class:
            filename_parts.append(selected_class)
        if selected_stream:
            filename_parts.append(selected_stream)
        if selected_gender:
            filename_parts.append(selected_gender)

    if selected_year:
        try:
            year_obj = AcademicYear.objects.get(id=selected_year)
            filename_parts.append(str(year_obj.year))
        except AcademicYear.DoesNotExist:
            pass

    filename = "_".join(filename_parts) + ".xlsx"

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = (
        f'attachment; filename="{filename}"'
    )

    wb.save(response)

    return response

@role_permission_required("add_student", "students")
def import_students(request):

    if request.method == "POST":

        uploaded_file = request.FILES.get(
            "student_file"
        )

        if not uploaded_file:

            messages.error(
                request,
                "Please select a CSV or Excel file."
            )

            return redirect(
                "import_students"
            )

        try:

            records, headers = (
                _read_student_import_file(
                    uploaded_file
                )
            )

        except Exception as exc:

            messages.error(
                request,
                str(exc)
            )

            return redirect(
                "import_students"
            )

        normalized_headers = {
            str(header).strip().lower()
            for header in headers
            if header
        }

        missing = [
            column
            for column in IMPORT_REQUIRED_COLUMNS
            if column not in normalized_headers
        ]

        if missing:

            messages.error(
                request,
                "Missing required columns: "
                + ", ".join(missing)
            )

            return redirect(
                "import_students"
            )

        errors = []
        prepared = []

        existing_admissions = set(
            Student.objects.values_list(
                "admission_no",
                flat=True
            )
        )

        file_admissions = set()

        for row_number, raw_row in enumerate(
            records,
            start=2
        ):

            row = {
                str(key).strip().lower(): value
                for key, value in raw_row.items()
                if key
            }

            first_name = _clean_import_value(
                row.get("first_name")
            )
            middle_name = _clean_import_value(
                row.get("middle_name")
            )
            last_name = _clean_import_value(
                row.get("last_name")
            )
            gender = _clean_import_value(
                row.get("gender")
            )
            class_name = _clean_import_value(
                row.get("class_name")
            ).upper()
            stream = _clean_import_value(
                row.get("stream")
            )
            parent_name = _clean_import_value(
                row.get("parent_name")
            )
            parent_phone = _clean_import_value(
                row.get("parent_phone")
            )
            admission_no = _clean_import_value(
                row.get("admission_no")
            ).upper()

            date_value = row.get(
                "date_of_birth"
            )

            # Excel may provide a real date object.
            if hasattr(date_value, "date"):
                date_of_birth = date_value.date()
            else:
                date_of_birth = _parse_import_date(
                    date_value
                )

            row_errors = []

            if not first_name:
                row_errors.append(
                    "First name is required"
                )

            if not last_name:
                row_errors.append(
                    "Last name is required"
                )

            if gender not in ["Male", "Female"]:
                row_errors.append(
                    "Gender must be Male or Female"
                )

            if not date_of_birth:
                row_errors.append(
                    "Invalid date of birth"
                )

            if class_name not in IMPORT_CLASSES:
                row_errors.append(
                    "Invalid class: " + class_name
                )

            if not stream:
                row_errors.append(
                    "Stream is required"
                )

            if not parent_name:
                row_errors.append(
                    "Parent name is required"
                )

            if not parent_phone:
                row_errors.append(
                    "Parent phone is required"
                )

            if admission_no:

                if admission_no in existing_admissions:
                    row_errors.append(
                        "Admission number already exists: "
                        + admission_no
                    )

                if admission_no in file_admissions:
                    row_errors.append(
                        "Duplicate admission number in file: "
                        + admission_no
                    )

                file_admissions.add(
                    admission_no
                )

            if row_errors:

                errors.append({
                    "row": row_number,
                    "errors": row_errors,
                })

            else:

                prepared.append({
                    "first_name": first_name,
                    "middle_name": middle_name,
                    "last_name": last_name,
                    "gender": gender,
                    "date_of_birth": date_of_birth,
                    "class_name": class_name,
                    "stream": stream,
                    "parent_name": parent_name,
                    "parent_phone": parent_phone,
                    "admission_no": admission_no,
                })

        # Never partially import a file.
        if errors:

            return render(
                request,
                "students/import_students.html",
                {
                    "errors": errors,
                    "total_rows": len(records),
                    "valid_rows": len(prepared),
                    "filename": uploaded_file.name,
                }
            )

        imported = 0

        try:

            with transaction.atomic():

                for data in prepared:

                    admission_no = data.pop(
                        "admission_no"
                    )

                    student = Student(
                        **data,
                        is_active=True
                    )

                    if admission_no:
                        student.admission_no = (
                            admission_no
                        )

                    student.save()

                    fee_ok, fee_result = apply_admission_fee_structure(
                        student
                    )

                    if not fee_ok:
                        raise Exception(
                            f"Fee structure could not be applied for "
                            f"{student.admission_no}: {fee_result}"
                        )

                    create_parent_account(
                        student
                    )

                    imported += 1

        except Exception as exc:

            messages.error(
                request,
                "Import failed. No records were imported. "
                + str(exc)
            )

            return redirect(
                "import_students"
            )

        messages.success(
            request,
            f"Successfully imported {imported} student(s)."
        )

        return redirect(
            "student_management"
        )

    return render(
        request,
        "students/import_students.html"
    )


def _get_import_classes():
    try:
        from timetable.models import TimetableClass

        classes = (
            TimetableClass.objects
            .filter(active=True)
            .exclude(name="")
            .values_list("name", flat=True)
            .distinct()
        )

        return {
            str(name).strip().upper()
            for name in classes
            if name
        }

    except Exception:
        return set()


IMPORT_CLASSES = _get_import_classes()


IMPORT_REQUIRED_COLUMNS = [
    "first_name",
    "last_name",
    "gender",
    "date_of_birth",
    "class_name",
    "stream",
    "parent_name",
    "parent_phone",
]


IMPORT_COLUMNS = [
    "first_name",
    "middle_name",
    "last_name",
    "gender",
    "date_of_birth",
    "class_name",
    "stream",
    "parent_name",
    "parent_phone",
    "admission_no",
]


@role_permission_required("view_student", "students")
def student_import_template(request):

    response = HttpResponse(
        content_type="text/csv"
    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; filename="student_import_template.csv"'
    )

    writer = csv.writer(response)

    writer.writerow(
        IMPORT_COLUMNS
    )

    writer.writerow([
        "John",
        "Kamau",
        "Otieno",
        "Male",
        "2015-05-20",
        "G5",
        "Blue",
        "Jane Otieno",
        "0712345678",
        "",
    ])

    return response



def apply_admission_fee_structure(student):
    """
    Automatically create the current term FeeRecord for a newly admitted
    learner using the active FeeStructure configured for the learner's class.
    """

    from django.db.models import Sum
    from fees.models import AcademicYear, Term, FeeRecord, FeeStructure

    # Get the active academic year and active term.
    academic_year = (
        AcademicYear.objects
        .filter(is_active=True)
        .order_by("-year")
        .first()
    )

    term = (
        Term.objects
        .filter(is_active=True)
        .order_by("order")
        .first()
    )

    if not academic_year or not term:
        return False, "No active academic year or term is configured."

    if not student.class_name:
        return False, "Learner has no class configured."

    # Total all active fee items configured for this learner's class.
    total = (
        FeeStructure.objects
        .filter(
            academic_year=academic_year,
            term=term,
            class_name=student.class_name,
            is_active=True,
        )
        .aggregate(total=Sum("amount"))
        ["total"]
    )

    if total is None:
        return False, (
            f"No active fee structure found for "
            f"{student.class_name} - {academic_year.year} - {term.name}."
        )

    # Never create duplicate fee records.
    fee_record = (
        FeeRecord.objects
        .filter(
            student=student,
            academic_year=academic_year,
            term=term,
        )
        .first()
    )

    if fee_record:
        return True, fee_record

    # Calculate opening balance from the learner's previous record
    # within the same academic year, if one exists.
    previous_record = (
        FeeRecord.objects
        .filter(
            student=student,
            academic_year=academic_year,
        )
        .exclude(term=term)
        .select_related("term")
        .order_by("-term__order", "-id")
        .first()
    )

    opening_balance = (
        previous_record.balance
        if previous_record
        else 0
    )

    fee_record = FeeRecord.objects.create(
        student=student,
        academic_year=academic_year,
        term=term,
        opening_balance=opening_balance,
        amount_charged=total,
        old_academic_year=str(academic_year.year),
        old_term=str(term.order),
    )

    return True, fee_record

@role_permission_required("add_student",
"students")
def add_student(request):

    if request.method == "POST":

        form = StudentForm(
            request.POST,
            request.FILES
        )

        if form.is_valid():

            student = form.save()

            try:
                fee_ok, fee_result = apply_admission_fee_structure(student)

                if fee_ok:
                    messages.success(
                        request,
                        f"Fee structure applied: KSh {fee_result.amount_charged:,.2f}"
                    )
                else:
                    messages.warning(
                        request,
                        f"Learner admitted, but fee structure was not applied: {fee_result}"
                    )

            except Exception as exc:
                messages.warning(
                    request,
                    f"Learner admitted, but fee setup was not completed: {exc}"
                )

            try:
                create_parent_account(student)

            except Exception as exc:
                messages.warning(
                    request,
                    f"Learner admitted, but parent account setup was not completed: {exc}"
                )

            messages.success(
                request,
                f"Learner added successfully — Admission Number: {student.admission_no}"
            )

            return redirect("student_list")

    else:

        form = StudentForm()

    return render(
        request,
        "students/add_student.html",
        {
            "form": form
        }
    )


@role_permission_required("view_student", "students")
def print_class_lists(request):
    from timetable.models import TimetableClass
    from academic.models import Stream
    from accounts.models import SchoolBranding
    from fees.models import AcademicYear

    selected_class = request.GET.get("class", "").strip()
    selected_stream = request.GET.get("stream", "").strip()

    classes = (
        Student.objects
        .filter(is_active=True)
        .exclude(class_name="")
        .values_list("class_name", flat=True)
        .distinct()
        .order_by("class_name")
    )

    streams = (
        Student.objects
        .filter(is_active=True)
        .exclude(stream="")
        .values_list("stream", flat=True)
        .distinct()
        .order_by("stream")
    )

    # Printing is allowed only after both Class and Stream
    # have been selected.
    if selected_class and selected_stream:
        students = (
            Student.objects
            .filter(
                is_active=True,
                class_name=selected_class,
                stream=selected_stream
            )
            .order_by(
                "last_name",
                "first_name",
                "middle_name",
                "admission_no"
            )
        )
    else:
        students = Student.objects.none()

    students = students.order_by(
        "class_name",
        "stream",
        "last_name",
        "first_name",
        "admission_no",
    )

    branding = (
        SchoolBranding.objects
        .filter(is_active=True)
        .first()
    )

    return render(
        request,
        "students/print_class_lists.html",
        {
            "students": students,
            "classes": classes,
            "streams": streams,
            "selected_class": selected_class,
            "selected_stream": selected_stream,
            "ready_to_print": bool(selected_class and selected_stream),
            "branding": branding,
        },
    )

@role_permission_required("view_student", "students")
def student_list(request):

    students = Student.objects.filter(
        is_active=True
    ).order_by(
        "admission_no"
    )

    return render(
        request,
        "students/student_list.html",
        {
            "students": students
        }
    )


@role_permission_required("view_student", "students")
def student_profile(request, id):

    student = get_object_or_404(
        Student,
        id=id
    )

    return render(
        request,
        "students/student_profile.html",
        {
            "student": student
        }
    )


@role_permission_required("change_student", "students")
def edit_student(request, id):

    student = get_object_or_404(
        Student,
        id=id
    )

    if request.method == "POST":

        form = StudentForm(
            request.POST,
            request.FILES,
            instance=student
        )

        if form.is_valid():

            student = form.save()

            # Make sure parent account remains linked
            create_parent_account(student)

            messages.success(request, "Details updated successfully")

            return redirect("student_list")

    else:

        form = StudentForm(
            instance=student
        )

    return render(
        request,
        "students/add_student.html",
        {
            "form": form,
            "student": student
        }
    )


@role_permission_required("delete_student", "students")
def delete_student(request, id):

    student = get_object_or_404(
        Student,
        id=id
    )

    if request.method == "POST":

        action = request.POST.get("action", "deactivate")

        if action == "permanent_delete":
            student.delete()

            messages.success(
                request,
                "Student record permanently deleted."
            )

        else:
            student.is_active = False
            student.save(update_fields=["is_active"])

            messages.success(
                request,
                f"{student.get_full_name()} has been deactivated."
            )

        return redirect("student_list")

    return render(
        request,
        "students/delete_student.html",
        {
            "student": student
        }
    )


# ============================================================
# ONLINE ADMISSION APPLICATIONS
# ============================================================

@role_permission_required("view_onlineapplication", "students")
def online_applications(request):

    from django.db.models import Q

    applications = (
        OnlineApplication.objects
        .select_related("admitted_student")
        .order_by("-application_date")
    )

    search = request.GET.get("search", "").strip()
    status = request.GET.get("status", "").strip()

    if search:
        applications = applications.filter(
            Q(application_no__icontains=search)
            | Q(first_name__icontains=search)
            | Q(middle_name__icontains=search)
            | Q(last_name__icontains=search)
            | Q(parent_name__icontains=search)
            | Q(parent_phone__icontains=search)
        )

    if status:
        applications = applications.filter(
            status=status
        )

    return render(
        request,
        "students/online_applications.html",
        {
            "applications": applications,
            "search": search,
            "selected_status": status,
            "pending_count": OnlineApplication.objects.filter(
                status="Pending"
            ).count(),
            "admitted_count": OnlineApplication.objects.filter(
                status="Admitted"
            ).count(),
            "rejected_count": OnlineApplication.objects.filter(
                status="Rejected"
            ).count(),
        }
    )


@role_permission_required("view_onlineapplication", "students")
def online_application_detail(request, id):

    application = get_object_or_404(
        OnlineApplication,
        id=id
    )

    return render(
        request,
        "students/online_application_detail.html",
        {
            "application": application
        }
    )


@role_permission_required("change_onlineapplication", "students")
def admit_online_application(request, id):

    application = get_object_or_404(
        OnlineApplication,
        id=id
    )

    if request.method != "POST":
        return redirect(
            "online_application_detail",
            id=id
        )

    if application.status == "Admitted":
        messages.warning(
            request,
            "This applicant has already been admitted."
        )
        return redirect(
            "online_application_detail",
            id=id
        )

    if application.status == "Rejected":
        messages.error(
            request,
            "A rejected application cannot be admitted."
        )
        return redirect(
            "online_application_detail",
            id=id
        )

    try:

        with transaction.atomic():

            student = Student.objects.create(
                first_name=application.first_name,
                middle_name=application.middle_name,
                last_name=application.last_name,
                gender=application.gender,
                date_of_birth=application.date_of_birth,
                class_name=application.requested_class,
                stream=application.requested_stream,
                parent_name=application.parent_name,
                parent_phone=application.parent_phone,
            )

            # --------------------------------------------------------
            # CREATE STUDENT PORTAL ACCOUNT
            # --------------------------------------------------------
            username = student.admission_no.lower()
            password = student.admission_no

            user = User.objects.filter(
                username=username
            ).first()

            if not user:
                user = User.objects.create_user(
                    username=username,
                    password=password
                )
            else:
                user.set_password(password)
                user.save(
                    update_fields=["password"]
                )

            # Find an existing profile already linked to this student,
            # or the profile belonging to the portal user.
            student_profile = (
                UserProfile.objects
                .filter(student=student)
                .first()
            )

            if student_profile is None:
                student_profile = (
                    UserProfile.objects
                    .filter(user=user)
                    .first()
                )

            if student_profile is None:
                student_profile = UserProfile.objects.create(
                    user=user,
                    role="STUDENT",
                    student=student,
                    phone=application.parent_phone,
                )
            else:
                # If this profile belongs to another user, use its user
                # for the student's portal account.
                user = student_profile.user

                student_profile.role = "STUDENT"
                student_profile.student = student
                student_profile.phone = application.parent_phone
                student_profile.is_active = True
                student_profile.save()

            # Automatically create/update parent account
            create_parent_account(student)

            # Automatically apply the configured fee structure
            fee_ok, fee_result = apply_admission_fee_structure(student)

            if not fee_ok:
                raise Exception(
                    f"Fee structure could not be applied: {fee_result}"
                )

            from django.utils import timezone

            application.status = "Admitted"
            application.admitted_student = student
            application.processed_at = timezone.now()

            application.save(
                update_fields=[
                    "status",
                    "admitted_student",
                    "processed_at",
                ]
            )

        messages.success(
            request,
            f"{application.get_full_name()} admitted successfully. "
            f"Admission Number: {student.admission_no}"
        )

    except Exception as exc:

        import traceback
        traceback.print_exc()

        messages.error(
            request,
            "Admission failed: " + str(exc)
        )

    return redirect("online_applications")


@role_permission_required("change_onlineapplication", "students")
def reject_online_application(request, id):

    application = get_object_or_404(
        OnlineApplication,
        id=id
    )

    if request.method != "POST":
        return redirect(
            "online_application_detail",
            id=id
        )

    if application.status == "Admitted":
        messages.error(
            request,
            "An admitted applicant cannot be rejected."
        )
        return redirect(
            "online_application_detail",
            id=id
        )

    from django.utils import timezone

    application.status = "Rejected"
    application.rejection_reason = (
        request.POST.get(
            "rejection_reason",
            ""
        ).strip()
    )
    application.processed_at = timezone.now()

    application.save(
        update_fields=[
            "status",
            "rejection_reason",
            "processed_at",
        ]
    )

    messages.success(
        request,
        f"{application.get_full_name()} has been rejected."
    )

    return redirect("online_applications")


# =========================================================
# PUBLIC ONLINE ADMISSION APPLICATION
# =========================================================

def online_application_form(request):
    from .forms import OnlineApplicationForm

    if request.method == "POST":
        form = OnlineApplicationForm(request.POST)

        if form.is_valid():
            application = form.save()

            return render(
                request,
                "students/online_application_success.html",
                {
                    "application": application,
                }
            )
    else:
        form = OnlineApplicationForm()

    return render(
        request,
        "students/online_application_form.html",
        {
            "form": form,
        }
    )


# ============================================================
# ONLINE APPLICATION STATUS
# ============================================================

def application_status(request):
    from .models import OnlineApplication
    from accounts.models import SchoolBranding

    application = None
    branding = (
        SchoolBranding.objects
        .filter(is_active=True)
        .first()
    )
    searched = False
    error = ""

    if request.method == "POST":
        reference = request.POST.get(
            "status_reference", ""
        ).strip().upper()

        searched = True

        if reference:
            application = (
                OnlineApplication.objects
                .select_related("admitted_student")
                .filter(status_reference=reference)
                .first()
            )

            if not application:
                error = (
                    "No application was found using that "
                    "status number. Please check the number "
                    "and try again."
                )
        else:
            error = "Please enter your application status number."

    return render(
        request,
        "students/application_status.html",
        {
            "application": application,
            "searched": searched,
            "error": error,
            "branding": branding,
        },
    )


# ============================================================
# BULK PROMOTION
# ============================================================

@role_permission_required("add_studentpromotionhistory", "students")
def bulk_promotion(request):

    from django.contrib import messages
    from django.db import transaction
    from django.shortcuts import redirect, render
    from fees.models import AcademicYear
    from timetable.models import TimetableClass
    from academic.models import Stream
    from .models import Student, StudentPromotionHistory

    years = (
        AcademicYear.objects
        .filter(is_active=True)
        .order_by("-year")
    )

    classes = (
        TimetableClass.objects
        .filter(active=True)
        .exclude(name="")
        .values_list("name", flat=True)
        .distinct()
        .order_by("name")
    )

    streams = (
        Stream.objects
        .filter(is_active=True)
        .exclude(name="")
        .values_list("name", flat=True)
        .distinct()
        .order_by("name")
    )

    selected_year = request.GET.get(
        "current_year",
        ""
    ).strip()

    selected_class = request.GET.get(
        "current_class",
        ""
    ).strip()

    selected_stream = request.GET.get(
        "current_stream",
        ""
    ).strip()

    selected_ids = request.POST.getlist(
        "student_ids"
    )

    students = Student.objects.all()

    if selected_year:
        try:
            students = students.filter(
                fee_records__academic_year_id=int(
                    selected_year
                )
            ).distinct()
        except ValueError:
            selected_year = ""

    if selected_class:
        students = students.filter(
            class_name=selected_class
        )

    if selected_stream:
        students = students.filter(
            stream=selected_stream
        )

    students = students.order_by(
        "class_name",
        "stream",
        "last_name",
        "first_name",
        "admission_no",
    )

    # ========================================================
    # PROCESS PROMOTION
    # ========================================================

    if request.method == "POST":

        target_year_id = request.POST.get(
            "target_year"
        )

        target_class = request.POST.get(
            "target_class",
            ""
        ).strip()

        target_stream = request.POST.get(
            "target_stream",
            ""
        ).strip()

        if not selected_ids:
            messages.error(
                request,
                "Please select at least one learner to promote."
            )

        elif not target_year_id:
            messages.error(
                request,
                "Please select the next academic year."
            )

        elif not target_class:
            messages.error(
                request,
                "Please select the next class."
            )

        else:

            try:
                target_year = AcademicYear.objects.get(
                    id=int(target_year_id),
                    is_active=True
                )

                # Validate target class against Admin configuration.
                if not TimetableClass.objects.filter(
                    active=True,
                    name=target_class
                ).exists():
                    raise ValueError(
                        "The selected next class is not active "
                        "in Admin configuration."
                    )

                # Validate target stream when supplied.
                if target_stream and not Stream.objects.filter(
                    is_active=True,
                    name=target_stream
                ).exists():
                    raise ValueError(
                        "The selected next stream is not active "
                        "in Admin configuration."
                    )

                learners = Student.objects.filter(
                    id__in=selected_ids
                ).order_by("id")

                promoted_count = 0
                skipped_count = 0

                with transaction.atomic():

                    for student in learners:

                        # Prevent duplicate promotion
                        # into the same academic year.
                        already_promoted = (
                            StudentPromotionHistory.objects
                            .filter(
                                student=student,
                                academic_year=target_year
                            )
                            .exists()
                        )

                        if already_promoted:
                            skipped_count += 1
                            continue

                        StudentPromotionHistory.objects.create(
                            student=student,
                            academic_year=target_year,
                            previous_class=student.class_name,
                            previous_stream=student.stream,
                            promoted_class=target_class,
                            promoted_stream=target_stream,
                        )

                        student.class_name = target_class
                        student.stream = target_stream
                        student.save(
                            update_fields=[
                                "class_name",
                                "stream",
                            ]
                        )

                        promoted_count += 1

                if promoted_count:
                    messages.success(
                        request,
                        f"{promoted_count} learner(s) promoted "
                        f"successfully to {target_class}"
                        + (
                            f" - {target_stream}"
                            if target_stream else ""
                        )
                        + f" for {target_year.year}."
                    )

                if skipped_count:
                    messages.warning(
                        request,
                        f"{skipped_count} learner(s) were skipped "
                        f"because they were already promoted for "
                        f"{target_year.year}."
                    )

                return redirect(
                    "students:bulk_promotion"
                )

            except Exception as exc:
                messages.error(
                    request,
                    "Promotion failed: " + str(exc)
                )

    return render(
        request,
        "students/bulk_promotion.html",
        {
            "students": students,
            "years": years,
            "classes": classes,
            "streams": streams,
            "selected_year": selected_year,
            "selected_class": selected_class,
            "selected_stream": selected_stream,
        }
    )


# ============================================================
# STUDENT PROMOTION HISTORY
# ============================================================

@role_permission_required("view_studentpromotionhistory", "students")
def student_promotion_history(request, id):

    from django.shortcuts import get_object_or_404, render
    from .models import Student, StudentPromotionHistory

    student = get_object_or_404(
        Student,
        id=id
    )

    history = (
        StudentPromotionHistory.objects
        .filter(student=student)
        .select_related("academic_year")
        .order_by(
            "-academic_year__year",
            "-promoted_at"
        )
    )

    return render(
        request,
        "students/student_promotion_history.html",
        {
            "student": student,
            "history": history,
        }
    )



