from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import render, redirect

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Protection
from openpyxl.worksheet.datavalidation import DataValidation

from students.models import Student
from .models import Subject, Assessment, AssessmentWindow
from .mark_permissions import (
    is_admin,
    can_enter_marks,
    can_access_class,
    allowed_subject_ids_for_class,
)


def _open_windows():
    """
    Only assessment windows that are currently allowed
    to receive marks are returned.

    OPEN          -> allowed
    CLOSING SOON  -> allowed
    REOPENED      -> allowed
    CLOSED        -> blocked
    DEADLINE      -> blocked
    """
    windows = (
        AssessmentWindow.objects
        .filter(is_active=True)
        .select_related(
            "academic_year",
            "term",
            "assessment_type",
        )
        .order_by("-deadline")
    )

    return [
        window
        for window in windows
        if window.recording_open()
    ]


def _learner_name(student):
    return " ".join(
        part for part in [
            getattr(student, "first_name", ""),
            getattr(student, "middle_name", ""),
            getattr(student, "last_name", ""),
        ]
        if part
    ).strip()


def _build_workbook(rows, title):
    wb = Workbook()
    ws = wb.active
    ws.title = "Marks Upload"

    headers = [
        "Admission No",
        "Learner Name",
        "Subject",
        "Mark",
    ]

    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append(row)

    # Lock first 3 columns and unlock Mark column
    for row in ws.iter_rows(min_row=2):
        for cell in row[:3]:
            cell.protection = Protection(locked=True)

        row[3].protection = Protection(locked=False)

    # Mark validation: 0 - 100
    validation = DataValidation(
        type="decimal",
        operator="between",
        formula1="0",
        formula2="100",
        allow_blank=True,
    )

    validation.error = "Enter a mark between 0 and 100."
    validation.errorTitle = "Invalid Mark"
    validation.prompt = "Enter a mark from 0 to 100."
    validation.promptTitle = "Mark Entry"

    ws.add_data_validation(validation)
    validation.add(f"D2:D{max(ws.max_row, 2)}")

    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 15

    # Protect sheet while allowing mark entry
    ws.protection.sheet = True
    ws.protection.password = "LUHAN_MARKS"

    # Instructions
    instructions = wb.create_sheet("Instructions")

    instruction_rows = [
        ["EXCEL MARK ENTRY"],
        [title],
        [""],
        ["Instructions"],
        ["1. Do not change Admission No, Learner Name or Subject."],
        ["2. Enter marks only in the Mark column."],
        ["3. Marks must be between 0 and 100."],
        ["4. Leave the Mark cell blank if no mark is available."],
        ["5. Save the completed Excel file before uploading."],
        ["6. Do not rename the worksheet."],
        ["7. Do not change the column headings."],
    ]

    for row in instruction_rows:
        instructions.append(row)

    instructions["A1"].font = Font(bold=True, size=14)
    instructions["A4"].font = Font(bold=True)
    instructions.column_dimensions["A"].width = 85

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    response["Content-Disposition"] = (
        f'attachment; filename="{title}.xlsx"'
    )

    return response


def _build_class_matrix_workbook(
    learners,
    subjects,
    existing_scores,
    title,
    request_user,
):
    """
    Whole-class matrix workbook.

    Learners = vertical rows.
    Subjects = horizontal columns.
    Only authorized subject cells are unlocked.
    """

    wb = Workbook()
    ws = wb.active
    ws.title = "Marks Upload"

    subjects = list(subjects)

    headers = [
        "Admission No",
        "Learner Name",
        "Stream",
    ] + [subject.name for subject in subjects]

    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for student in learners:

        student_stream = getattr(
            student,
            "stream",
            None,
        )

        row = [
            student.admission_no,
            _learner_name(student),
            student_stream or "",
        ]

        for subject in subjects:

            authorized = can_enter_marks(
                request_user,
                student.class_name,
                student_stream,
                subject,
            )

            if authorized:
                value = existing_scores.get(
                    (student.pk, subject.pk),
                    "",
                )
            else:
                # Unauthorized subject cells remain blank
                # and locked.
                value = ""

            row.append(value)

        ws.append(row)

    # --------------------------------------------------------
    # PROTECTION
    # --------------------------------------------------------

    for row in ws.iter_rows(min_row=2):

        # Identity columns
        for cell in row[:3]:
            cell.protection = Protection(locked=True)

        # Subject mark columns
        for index, subject in enumerate(
            subjects,
            start=3,
        ):

            cell = row[index]

            student_row_number = cell.row
            admission_no = ws.cell(
                student_row_number,
                1,
            ).value

            try:
                student = next(
                    student
                    for student in learners
                    if str(student.admission_no).strip()
                    == str(admission_no).strip()
                )
            except StopIteration:
                cell.protection = Protection(locked=True)
                continue

            student_stream = getattr(
                student,
                "stream",
                None,
            )

            if can_enter_marks(
                request_user,
                student.class_name,
                student_stream,
                subject,
            ):
                cell.protection = Protection(locked=False)
            else:
                cell.protection = Protection(locked=True)

    # --------------------------------------------------------
    # MARK VALIDATION
    # --------------------------------------------------------

    validation = DataValidation(
        type="decimal",
        operator="between",
        formula1="0",
        formula2="100",
        allow_blank=True,
    )

    validation.error = "Enter a mark between 0 and 100."
    validation.errorTitle = "Invalid Mark"
    validation.prompt = "Enter a mark from 0 to 100."
    validation.promptTitle = "Mark Entry"

    ws.add_data_validation(validation)

    if ws.max_row >= 2 and len(subjects) > 0:

        first_subject_column = 4
        last_subject_column = 3 + len(subjects)

        from openpyxl.utils import get_column_letter

        validation.add(
            f"{get_column_letter(first_subject_column)}2:"
            f"{get_column_letter(last_subject_column)}{ws.max_row}"
        )

    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 16

    from openpyxl.utils import get_column_letter

    for index in range(4, 4 + len(subjects)):
        ws.column_dimensions[
            get_column_letter(index)
        ].width = 18

    ws.protection.sheet = True
    ws.protection.password = "LUHAN_MARKS"

    # --------------------------------------------------------
    # INSTRUCTIONS
    # --------------------------------------------------------

    instructions = wb.create_sheet("Instructions")

    instruction_rows = [
        ["WHOLE CLASS EXCEL MARK ENTRY"],
        [title],
        [""],
        ["Format"],
        ["Learners are listed vertically, one learner per row."],
        ["Subjects are listed horizontally across the top."],
        [""],
        ["Instructions"],
        ["1. Do not change Admission No, Learner Name or Stream."],
        ["2. Enter marks only in unlocked subject cells."],
        ["3. Marks must be between 0 and 100."],
        ["4. Leave a mark blank if it is not available."],
        ["5. Do not rename the Marks Upload worksheet."],
        ["6. Do not change the subject column headings."],
        ["7. Save the completed Excel file before uploading."],
        ["8. Only subjects assigned to you can be entered."],
    ]

    for row in instruction_rows:
        instructions.append(row)

    instructions["A1"].font = Font(
        bold=True,
        size=14,
    )

    instructions["A4"].font = Font(bold=True)
    instructions["A8"].font = Font(bold=True)
    instructions.column_dimensions["A"].width = 90

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    response["Content-Disposition"] = (
        f'attachment; filename="{title}.xlsx"'
    )

    return response


@login_required
def individual_excel_marks(request):

    windows = _open_windows()

    selected_window_id = request.GET.get("window")
    selected_student_id = request.GET.get("student")

    # --------------------------------------------------------
    # SECURITY-FIRST STUDENT LIST
    # --------------------------------------------------------

    all_students = Student.objects.all().order_by(
        "class_name",
        "stream",
        "first_name",
        "last_name",
    )

    if is_admin(request.user):

        students = all_students

    else:

        allowed_student_ids = []

        for student in all_students:

            if can_access_class(
                request.user,
                student.class_name,
                getattr(student, "stream", None),
            ):
                allowed_student_ids.append(student.id)

        students = all_students.filter(
            id__in=allowed_student_ids
        )

    # --------------------------------------------------------
    # GENERATE INDIVIDUAL EXCEL
    # --------------------------------------------------------

    if selected_window_id and selected_student_id:

        window = next(
            (
                w for w in windows
                if str(w.id) == str(selected_window_id)
            ),
            None,
        )

        if not window:

            messages.error(
                request,
                "This assessment is not currently open for Excel mark entry."
            )

            return redirect(
                "academic:excel_individual"
            )

        try:

            student = all_students.get(
                id=selected_student_id
            )

        except Student.DoesNotExist:

            messages.error(
                request,
                "The selected learner could not be found."
            )

            return redirect(
                "academic:excel_individual"
            )

        # ----------------------------------------------------
        # CLASS ACCESS SECURITY
        # ----------------------------------------------------

        student_stream = getattr(
            student,
            "stream",
            None,
        )

        if not can_access_class(
            request.user,
            student.class_name,
            student_stream,
        ):

            messages.error(
                request,
                "You are not authorized to generate marks "
                "for this learner's class."
            )

            return redirect(
                "academic:excel_individual"
            )

        # ----------------------------------------------------
        # SUBJECT SECURITY
        #
        # Admin       -> all subjects
        # Class teacher -> all subjects in assigned class
        # Teacher     -> assigned subjects only
        # ----------------------------------------------------

        allowed_subject_ids = allowed_subject_ids_for_class(
            request.user,
            student.class_name,
            student_stream,
        )

        if allowed_subject_ids is None:

            subjects = Subject.objects.filter(
                is_active=True
            ).order_by("name")

        else:

            subjects = Subject.objects.filter(
                id__in=allowed_subject_ids,
                is_active=True,
            ).order_by("name")

        if not subjects.exists():

            messages.error(
                request,
                "You are not authorized to enter marks "
                "for any subject in this learner's class."
            )

            return redirect(
                "academic:excel_individual"
            )

        # ----------------------------------------------------
        # LOAD EXISTING MARKS FROM THE EXACT ASSESSMENT WINDOW
        # ----------------------------------------------------

        existing_scores = {
            (item["student_id"], item["subject_id"]): item["score"]
            for item in Assessment.objects.filter(
                student_id=student.pk,
                academic_year_id=window.academic_year_id,
                term_id=window.term_id,
                assessment_type_id=window.assessment_type_id,
                subject_id__in=subjects.values_list("id", flat=True),
            ).values(
                "student_id",
                "subject_id",
                "score",
            )
        }

        rows = []

        learner_name = _learner_name(student)

        for subject in subjects:

            # Final server-side permission check.
            if not can_enter_marks(
                request.user,
                student.class_name,
                student_stream,
                subject,
            ):
                continue

            rows.append([
                student.admission_no,
                learner_name,
                subject.name,
                existing_scores.get(
                    (student.pk, subject.pk),
                    "",
                ),
            ])

        if not rows:

            messages.error(
                request,
                "No subjects are available for your mark-entry permission."
            )

            return redirect(
                "academic:excel_individual"
            )

        safe_name = (
            learner_name
            .replace("/", "-")
            .replace("\\", "-")
        )

        safe_assessment = (
            window.assessment_type.name
            .replace("/", "-")
            .replace("\\", "-")
        )

        filename_title = (
            f"Individual_{safe_name}_{safe_assessment}_Marks"
        )

        return _build_workbook(
            rows,
            filename_title,
        )

    return render(
        request,
        "academic/excel_marks_individual.html",
        {
            "windows": windows,
            "students": students,
        },
    )


@login_required
def class_excel_marks(request):

    windows = _open_windows()

    selected_window_id = request.GET.get("window")
    selected_class = request.GET.get(
        "class_name",
        "",
    ).strip()

    selected_stream = request.GET.get(
        "stream",
        "",
    ).strip()

    # --------------------------------------------------------
    # CLASS LIST SECURITY
    # --------------------------------------------------------

    all_class_names = (
        Student.objects
        .exclude(
            class_name__isnull=True
        )
        .exclude(
            class_name__exact=""
        )
        .values_list(
            "class_name",
            flat=True,
        )
        .distinct()
        .order_by("class_name")
    )

    if is_admin(request.user):

        class_names = all_class_names

    else:

        allowed_classes = set()

        for class_name in all_class_names:

            if can_access_class(
                request.user,
                class_name,
                None,
            ):

                allowed_classes.add(
                    str(class_name)
                )

        class_names = [
            class_name
            for class_name in all_class_names
            if str(class_name) in allowed_classes
        ]

    streams = (
        Student.objects
        .values_list(
            "stream",
            flat=True,
        )
        .exclude(
            stream__isnull=True
        )
        .exclude(
            stream__exact=""
        )
        .distinct()
        .order_by("stream")
    )

    # --------------------------------------------------------
    # GENERATE CLASS EXCEL
    # --------------------------------------------------------

    if selected_window_id and selected_class:

        window = next(
            (
                w for w in windows
                if str(w.id) == str(selected_window_id)
            ),
            None,
        )

        if not window:

            messages.error(
                request,
                "This assessment is not currently open for Excel mark entry."
            )

            return redirect(
                "academic:excel_class"
            )

        # ----------------------------------------------------
        # CLASS ACCESS SECURITY
        # ----------------------------------------------------

        if not can_access_class(
            request.user,
            selected_class,
            selected_stream or None,
        ):

            messages.error(
                request,
                "You are not authorized to generate marks "
                "for this class."
            )

            return redirect(
                "academic:excel_class"
            )

        learners = Student.objects.filter(
            class_name=selected_class
        ).order_by(
            "stream",
            "first_name",
            "last_name",
        )

        if selected_stream:

            learners = learners.filter(
                stream=selected_stream
            )

        if not learners.exists():

            messages.error(
                request,
                "No learners were found for the selected class and stream."
            )

            return redirect(
                "academic:excel_class"
            )

        # ----------------------------------------------------
        # DETERMINE SUBJECTS FOR THE CLASS
        # ----------------------------------------------------

        allowed_subject_ids = allowed_subject_ids_for_class(
            request.user,
            selected_class,
            selected_stream or None,
        )

        if allowed_subject_ids is None:

            subjects = Subject.objects.filter(
                is_active=True
            ).order_by("name")

        else:

            subjects = Subject.objects.filter(
                id__in=allowed_subject_ids,
                is_active=True,
            ).order_by("name")

        if not subjects.exists():

            messages.error(
                request,
                "You are not authorized to enter marks "
                "for any subject in this class."
            )

            return redirect(
                "academic:excel_class"
            )

        # ----------------------------------------------------
        # LOAD EXISTING MARKS FROM THE EXACT ASSESSMENT WINDOW
        # ----------------------------------------------------

        existing_scores = {
            (item["student_id"], item["subject_id"]): item["score"]
            for item in Assessment.objects.filter(
                student_id__in=learners.values_list("id", flat=True),
                academic_year_id=window.academic_year_id,
                term_id=window.term_id,
                assessment_type_id=window.assessment_type_id,
                subject_id__in=subjects.values_list("id", flat=True),
            ).values(
                "student_id",
                "subject_id",
                "score",
            )
        }

        # ----------------------------------------------------
        # BUILD WHOLE-CLASS MATRIX
        #
        # Learners = rows
        # Subjects = columns
        # ----------------------------------------------------

        authorized_combinations = 0

        for student in learners:

            student_stream = getattr(
                student,
                "stream",
                None,
            )

            for subject in subjects:

                if can_enter_marks(
                    request.user,
                    student.class_name,
                    student_stream,
                    subject,
                ):
                    authorized_combinations += 1

        if not authorized_combinations:

            messages.error(
                request,
                "No authorized learner-subject combinations "
                "were found for this class."
            )

            return redirect(
                "academic:excel_class"
            )

        # ----------------------------------------------------
        # FILE NAME
        # ----------------------------------------------------

        safe_class = (
            str(selected_class)
            .replace("/", "-")
            .replace("\\", "-")
            .replace(" ", "_")
        )

        safe_stream = (
            str(selected_stream)
            .replace("/", "-")
            .replace("\\", "-")
            .replace(" ", "_")
            if selected_stream
            else ""
        )

        safe_assessment = (
            window.assessment_type.name
            .replace("/", "-")
            .replace("\\", "-")
            .replace(" ", "_")
        )

        filename_parts = [
            "Class",
            safe_class,
        ]

        if safe_stream:
            filename_parts.append(
                safe_stream
            )

        filename_parts.extend([
            safe_assessment,
            "Marks",
        ])

        filename_title = "_".join(
            filename_parts
        )

        return _build_class_matrix_workbook(
            learners,
            subjects,
            existing_scores,
            filename_title,
            request.user,
        )

    return render(
        request,
        "academic/excel_marks_class.html",
        {
            "windows": windows,
            "class_names": class_names,
            "streams": streams,
        },
    )

