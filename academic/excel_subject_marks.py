
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import render, redirect

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Protection
from openpyxl.worksheet.datavalidation import DataValidation

from students.models import Student
from .models import Subject, Assessment, AssessmentWindow
from .mark_permissions import (
    is_admin,
    can_access_class,
    can_enter_marks,
    allowed_subject_ids_for_class,
)


def _learner_name(student):
    return " ".join(
        part for part in [
            getattr(student, "first_name", ""),
            getattr(student, "middle_name", ""),
            getattr(student, "last_name", ""),
        ]
        if part
    ).strip()


def _open_windows():
    return list(
        AssessmentWindow.objects.filter(
            is_active=True
        ).select_related(
            "academic_year",
            "term",
            "assessment_type",
        ).order_by("-deadline")
    )


def _classes_for_user(user):
    classes = []

    for class_name in (
        Student.objects.filter(
            class_name__isnull=False
        )
        .exclude(class_name="")
        .values_list("class_name", flat=True)
        .distinct()
        .order_by("class_name")
    ):
        if is_admin(user) or can_access_class(user, class_name, None):
            classes.append(class_name)

    return classes


def _subjects_for_class(user, class_name, stream):
    """
    Return subjects the user may enter for a class.

    If a stream is selected, permissions are checked against
    that specific stream.

    If no stream is selected, combine subjects from all streams
    assigned to the teacher in that class. This allows the
    Subject Excel pages to populate correctly before a stream
    is selected.
    """

    if not class_name:
        if is_admin(user):
            return Subject.objects.filter(
                is_active=True
            ).order_by("name")

        return Subject.objects.none()

    # Admins can access every active subject.
    if is_admin(user):
        return Subject.objects.filter(
            is_active=True
        ).order_by("name")

    # Specific stream selected.
    if stream:
        allowed = allowed_subject_ids_for_class(
            user,
            class_name,
            stream,
        )

        if allowed is None:
            return Subject.objects.filter(
                is_active=True
            ).order_by("name")

        return Subject.objects.filter(
            id__in=allowed,
            is_active=True,
        ).order_by("name")

    # --------------------------------------------------------
    # No stream selected:
    # collect subjects allowed in ANY assigned stream
    # of this class.
    # --------------------------------------------------------

    class_streams = list(
        Student.objects.filter(
            class_name__iexact=class_name,
            stream__isnull=False,
        )
        .exclude(stream="")
        .values_list("stream", flat=True)
        .distinct()
    )

    allowed_ids = set()

    # Also check the class itself where assignments have
    # no stream specified.
    allowed = allowed_subject_ids_for_class(
        user,
        class_name,
        None,
    )

    if allowed is None:
        return Subject.objects.filter(
            is_active=True
        ).order_by("name")

    allowed_ids.update(allowed)

    for class_stream in class_streams:
        stream_allowed = allowed_subject_ids_for_class(
            user,
            class_name,
            class_stream,
        )

        if stream_allowed is None:
            return Subject.objects.filter(
                is_active=True
            ).order_by("name")

        allowed_ids.update(stream_allowed)

    return Subject.objects.filter(
        id__in=allowed_ids,
        is_active=True,
    ).order_by("name")


# ============================================================
# GENERATE SUBJECT MARKS EXCEL
# ============================================================

@login_required
def generate_subject_marks(request):

    windows = _open_windows()
    classes = _classes_for_user(request.user)

    selected_window_id = request.GET.get("window") or request.POST.get("window")
    selected_class = request.GET.get("class_name") or request.POST.get("class_name") or ""
    selected_stream = request.GET.get("stream") or request.POST.get("stream") or ""
    selected_subject_id = request.GET.get("subject") or request.POST.get("subject")

    subjects = _subjects_for_class(
        request.user,
        selected_class,
        selected_stream,
    )

    streams = []

    if selected_class:
        streams = list(
            Student.objects.filter(
                class_name__iexact=selected_class,
                stream__isnull=False,
            )
            .exclude(stream="")
            .values_list("stream", flat=True)
            .distinct()
            .order_by("stream")
        )

    if request.method == "POST":

        if not selected_window_id:
            messages.error(request, "Please select an assessment.")
            return redirect("academic:excel_generate_subject")

        if not selected_class:
            messages.error(request, "Please select a class.")
            return redirect("academic:excel_generate_subject")

        if not selected_subject_id:
            messages.error(request, "Please select a subject.")
            return redirect("academic:excel_generate_subject")

        try:
            window = AssessmentWindow.objects.select_related(
                "academic_year",
                "term",
                "assessment_type",
            ).get(
                pk=selected_window_id,
                is_active=True,
            )
        except AssessmentWindow.DoesNotExist:
            messages.error(request, "The selected assessment is not active.")
            return redirect("academic:excel_generate_subject")

        if not window.recording_open():
            messages.error(
                request,
                "This assessment is closed. Marks cannot be generated for entry."
            )
            return redirect("academic:excel_generate_subject")

        try:
            subject = Subject.objects.get(
                pk=selected_subject_id,
                is_active=True,
            )
        except Subject.DoesNotExist:
            messages.error(request, "Invalid subject.")
            return redirect("academic:excel_generate_subject")

        if not can_access_class(
            request.user,
            selected_class,
            selected_stream or None,
        ):
            messages.error(
                request,
                "You are not authorized to access this class."
            )
            return redirect("academic:excel_generate_subject")

        if not can_enter_marks(
            request.user,
            selected_class,
            selected_stream or None,
            subject,
        ):
            messages.error(
                request,
                "You are not authorized to enter marks for this subject in this class."
            )
            return redirect("academic:excel_generate_subject")

        learners = Student.objects.filter(
            class_name__iexact=selected_class,
        ).order_by(
            "admission_no",
            "first_name",
            "last_name",
        )

        if selected_stream:
            learners = learners.filter(
                stream__iexact=selected_stream
            )

        score_map = {
            item["student_id"]: item["score"]
            for item in Assessment.objects.filter(
                student_id__in=learners.values_list("id", flat=True),
                academic_year_id=window.academic_year_id,
                term_id=window.term_id,
                assessment_type_id=window.assessment_type_id,
                subject_id=subject.id,
            ).values(
                "student_id",
                "score",
            )
        }

        wb = Workbook()
        ws = wb.active
        ws.title = "Marks Upload"

        headers = [
            "Admission No",
            "Learner Name",
            "Class",
            "Stream",
            "Subject",
            "Mark",
        ]

        ws.append(headers)

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        for student in learners:

            if not can_enter_marks(
                request.user,
                student.class_name,
                getattr(student, "stream", None),
                subject,
            ):
                continue

            ws.append([
                student.admission_no,
                _learner_name(student),
                student.class_name,
                getattr(student, "stream", "") or "",
                subject.name,
                score_map.get(student.pk, ""),
            ])

        # Lock everything except Mark
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.protection = Protection(locked=True)

            row[5].protection = Protection(locked=False)

        validation = DataValidation(
            type="decimal",
            operator="between",
            formula1="0",
            formula2="100",
            allow_blank=True,
        )

        validation.error = "Enter a mark between 0 and 100."
        validation.errorTitle = "Invalid mark"
        validation.prompt = "Enter a mark from 0 to 100."
        validation.promptTitle = "Mark"

        ws.add_data_validation(validation)
        validation.add(f"F2:F{max(ws.max_row, 2)}")

        ws.freeze_panes = "F2"
        ws.auto_filter.ref = ws.dimensions

        widths = {
            "A": 18,
            "B": 32,
            "C": 15,
            "D": 15,
            "E": 28,
            "F": 12,
        }

        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        ws.protection.sheet = True
        ws.protection.password = "LUHAN_MARKS"

        info = wb.create_sheet("Instructions")

        info.append(["ASSESSMENT", str(window.assessment_type)])
        info.append(["ACADEMIC YEAR", str(window.academic_year)])
        info.append(["TERM", str(window.term)])
        info.append(["SUBJECT", subject.name])
        info.append(["CLASS", selected_class])
        info.append(["STREAM", selected_stream or "All"])
        info.append([])
        info.append(["INSTRUCTIONS"])
        info.append(["Only edit the Mark column."])
        info.append(["Marks must be between 0 and 100."])
        info.append(["Do not change Admission No, Learner Name, Class, Stream or Subject."])
        info.append(["Upload this file using Upload Subject Marks Excel."])

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        filename = (
            f"{selected_class}"
            f"{'_' + selected_stream if selected_stream else ''}_"
            f"{subject.name}_"
            f"{window.assessment_type}_Marks.xlsx"
        )

        response = HttpResponse(
            output.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

        response["Content-Disposition"] = (
            f'attachment; filename="{filename}"'
        )

        return response

    return render(
        request,
        "academic/excel_subject_generate.html",
        {
            "windows": windows,
            "classes": classes,
            "streams": streams,
            "subjects": subjects,
            "selected_window_id": selected_window_id,
            "selected_class": selected_class,
            "selected_stream": selected_stream,
            "selected_subject_id": selected_subject_id,
        },
    )


# ============================================================
# UPLOAD SUBJECT MARKS EXCEL
# ============================================================

@login_required
def upload_subject_marks(request):

    windows = _open_windows()
    classes = _classes_for_user(request.user)

    selected_window_id = request.POST.get("window") or request.GET.get("window")
    selected_class = request.POST.get("class_name") or request.GET.get("class_name") or ""
    selected_stream = request.POST.get("stream") or request.GET.get("stream") or ""
    selected_subject_id = request.POST.get("subject") or request.GET.get("subject")

    subjects = _subjects_for_class(
        request.user,
        selected_class,
        selected_stream,
    )

    streams = []

    if selected_class:
        streams = list(
            Student.objects.filter(
                class_name__iexact=selected_class,
                stream__isnull=False,
            )
            .exclude(stream="")
            .values_list("stream", flat=True)
            .distinct()
            .order_by("stream")
        )

    if request.method == "POST":

        if not selected_window_id:
            messages.error(request, "Please select an assessment.")
            return redirect("academic:excel_upload_subject")

        if not selected_class:
            messages.error(request, "Please select a class.")
            return redirect("academic:excel_upload_subject")

        if not selected_subject_id:
            messages.error(request, "Please select a subject.")
            return redirect("academic:excel_upload_subject")

        uploaded_file = request.FILES.get("excel_file")

        if not uploaded_file:
            messages.error(request, "Please select an Excel file.")
            return redirect("academic:excel_upload_subject")

        try:
            window = AssessmentWindow.objects.select_related(
                "academic_year",
                "term",
                "assessment_type",
            ).get(
                pk=selected_window_id,
                is_active=True,
            )
        except AssessmentWindow.DoesNotExist:
            messages.error(request, "The selected assessment is not active.")
            return redirect("academic:excel_upload_subject")

        if not window.recording_open():
            messages.error(
                request,
                "This assessment is closed. Marks cannot be uploaded."
            )
            return redirect("academic:excel_upload_subject")

        try:
            subject = Subject.objects.get(
                pk=selected_subject_id,
                is_active=True,
            )
        except Subject.DoesNotExist:
            messages.error(request, "Invalid subject.")
            return redirect("academic:excel_upload_subject")

        if not can_access_class(
            request.user,
            selected_class,
            selected_stream or None,
        ):
            messages.error(
                request,
                "You are not authorized to access this class."
            )
            return redirect("academic:excel_upload_subject")

        if not can_enter_marks(
            request.user,
            selected_class,
            selected_stream or None,
            subject,
        ):
            messages.error(
                request,
                "You are not authorized to enter marks for this subject in this class."
            )
            return redirect("academic:excel_upload_subject")

        try:
            wb = load_workbook(
                uploaded_file,
                data_only=True,
            )
            ws = wb.active
        except Exception:
            messages.error(
                request,
                "The uploaded file is not a valid Excel workbook."
            )
            return redirect("academic:excel_upload_subject")

        headers = [
            cell.value
            for cell in ws[1]
        ]

        normalized_headers = [
            str(value).strip().casefold()
            if value is not None else ""
            for value in headers
        ]

        required = [
            "admission no",
            "learner name",
            "class",
            "stream",
            "subject",
            "mark",
        ]

        missing = [
            h for h in required
            if h not in normalized_headers
        ]

        if missing:
            messages.error(
                request,
                "Invalid Excel format. Missing: " + ", ".join(missing)
            )
            return redirect("academic:excel_upload_subject")

        col = {
            name: normalized_headers.index(name)
            for name in required
        }

        valid_rows = []
        errors = []
        seen = set()

        for row_number, row in enumerate(
            ws.iter_rows(min_row=2, values_only=True),
            start=2,
        ):

            admission_no = row[col["admission no"]]
            excel_class = row[col["class"]]
            excel_stream = row[col["stream"]]
            excel_subject = row[col["subject"]]
            raw_mark = row[col["mark"]]

            if not admission_no:
                continue

            admission_no = str(admission_no).strip()

            if str(excel_class or "").strip().casefold() != selected_class.strip().casefold():
                errors.append(
                    f"Row {row_number}: learner belongs to a different class."
                )
                continue

            if selected_stream:
                if str(excel_stream or "").strip().casefold() != selected_stream.strip().casefold():
                    errors.append(
                        f"Row {row_number}: learner belongs to a different stream."
                    )
                    continue

            if str(excel_subject or "").strip().casefold() != subject.name.strip().casefold():
                errors.append(
                    f"Row {row_number}: wrong subject."
                )
                continue

            try:
                student = Student.objects.get(
                    admission_no=admission_no
                )
            except Student.DoesNotExist:
                errors.append(
                    f"Row {row_number}: learner {admission_no} was not found."
                )
                continue

            if not can_enter_marks(
                request.user,
                student.class_name,
                getattr(student, "stream", None),
                subject,
            ):
                errors.append(
                    f"Row {row_number}: you are not authorized for {admission_no}."
                )
                continue

            if raw_mark in (None, ""):
                errors.append(
                    f"Row {row_number}: mark is blank."
                )
                continue

            try:
                numeric_mark = float(raw_mark)
            except (TypeError, ValueError):
                errors.append(
                    f"Row {row_number}: invalid mark."
                )
                continue

            if numeric_mark < 0 or numeric_mark > 100:
                errors.append(
                    f"Row {row_number}: mark must be between 0 and 100."
                )
                continue

            key = student.pk

            if key in seen:
                errors.append(
                    f"Row {row_number}: duplicate learner {admission_no}."
                )
                continue

            seen.add(key)

            valid_rows.append({
                "student": student,
                "score": numeric_mark,
            })

        if errors:
            messages.error(
                request,
                "Upload rejected. No marks were saved."
            )

            for error in errors[:10]:
                messages.error(request, error)

            return redirect("academic:excel_upload_subject")

        for item in valid_rows:
            Assessment.objects.update_or_create(
                student=item["student"],
                academic_year_id=window.academic_year_id,
                term_id=window.term_id,
                subject=subject,
                assessment_type_id=window.assessment_type_id,
                defaults={
                    "score": item["score"],
                },
            )

        messages.success(
            request,
            f"{len(valid_rows)} {subject.name} marks uploaded successfully."
        )

        return redirect("academic:excel_upload_subject")

    return render(
        request,
        "academic/excel_subject_upload.html",
        {
            "windows": windows,
            "classes": classes,
            "streams": streams,
            "subjects": subjects,
            "selected_window_id": selected_window_id,
            "selected_class": selected_class,
            "selected_stream": selected_stream,
            "selected_subject_id": selected_subject_id,
        },
    )
