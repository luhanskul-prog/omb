from openpyxl import load_workbook

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.shortcuts import redirect, render

from students.models import Student
from timetable.legacy_compat import (
    TeacherClassAssignment,
    TeacherTeachingAssignment,
)

from .models import (
    AcademicYear,
    Term,
    Subject,
    AssessmentType,
    Assessment,
    AssessmentWindow,
)

from .mark_permissions import (
    is_admin,
    can_enter_marks,
    get_scheduling_teacher,
    _same_class,
    _same_stream,
)


def _allowed_class_streams_for_upload(user):
    """
    Return the class/stream combinations the user may use
    for Excel mark upload.

    Admin:
        all active learner class/stream combinations.

    Class teacher:
        their assigned class/stream combinations.

    Subject teacher:
        their teaching class/stream combinations.
    """

    combinations = set()

    if is_admin(user):
        rows = (
            Student.objects
            .exclude(class_name__isnull=True)
            .exclude(class_name="")
            .values_list("class_name", "stream")
            .distinct()
        )

        for class_name, stream in rows:
            combinations.add(
                (
                    str(class_name).strip(),
                    str(stream or "").strip(),
                )
            )

        return combinations

    teacher = get_scheduling_teacher(user)

    if not teacher:
        return combinations

    class_rows = TeacherClassAssignment.objects.filter(
        teacher=teacher,
        is_active=True,
    ).values_list(
        "class_name",
        "stream",
    )

    teaching_rows = TeacherTeachingAssignment.objects.filter(
        teacher=teacher,
        is_active=True,
    ).values_list(
        "class_name",
        "stream",
    )

    for class_name, stream in list(class_rows) + list(teaching_rows):
        combinations.add(
            (
                str(class_name).strip(),
                str(stream or "").strip(),
            )
        )

    return combinations


def _allowed_subject_ids_for_upload(
    user,
    class_name,
    stream,
):
    """
    Return academic.Subject IDs permitted for this exact
    class/stream.

    Admin:
        all active subjects.

    Class teacher:
        all active subjects.

    Subject teacher:
        only subjects assigned to that class/stream.

    timetable.TimetableSubject and academic.Subject are separate
    models, therefore matching is done by subject CODE.
    """

    if not class_name:
        return set()

    if is_admin(user):
        return set(
            Subject.objects
            .filter(is_active=True)
            .values_list("id", flat=True)
        )

    teacher = get_scheduling_teacher(user)

    if not teacher:
        return set()

    # First determine whether this teacher is a class teacher
    # for this exact class/stream.
    is_class_teacher = False

    for assignment in TeacherClassAssignment.objects.filter(
        teacher=teacher,
        is_active=True,
    ):
        if _same_class(
            class_name,
            assignment.class_name,
        ) and _same_stream(
            stream,
            assignment.stream,
        ):
            is_class_teacher = True
            break

    if is_class_teacher:
        return set(
            Subject.objects
            .filter(is_active=True)
            .values_list("id", flat=True)
        )

    # Subject teacher: collect scheduling subject codes.
    allowed_codes = set()

    assignments = (
        TeacherTeachingAssignment.objects
        .filter(
            teacher=teacher,
            is_active=True,
        )
        .select_related("subject")
    )

    for assignment in assignments:

        if not _same_class(
            class_name,
            assignment.class_name,
        ):
            continue

        if not _same_stream(
            stream,
            assignment.stream,
        ):
            continue

        code = getattr(
            assignment.subject,
            "code",
            None,
        )

        if code:
            allowed_codes.add(
                str(code).strip().casefold()
            )

    allowed_ids = set()

    for subject in Subject.objects.filter(
        is_active=True
    ):
        code = getattr(
            subject,
            "code",
            None,
        )

        if (
            code
            and str(code).strip().casefold()
            in allowed_codes
        ):
            allowed_ids.add(subject.id)

    return allowed_ids


@login_required
def upload_marks_excel(request):

    open_windows = []

    for window in AssessmentWindow.objects.filter(
        is_active=True
    ).select_related(
        "academic_year",
        "term",
        "assessment_type",
    ):
        if window.recording_open():
            open_windows.append(window)

    # Assessment & Marks Management is the source of truth.
    # The selected AssessmentWindow supplies the academic year,
    # term and assessment type.

    # --------------------------------------------------------
    # PERMITTED CLASS/STREAM COMBINATIONS
    # --------------------------------------------------------

    allowed_combinations = (
        _allowed_class_streams_for_upload(request.user)
    )

    classes = sorted(
        {
            class_name
            for class_name, stream in allowed_combinations
        },
        key=str.casefold,
    )

    selected_class = (
        request.GET.get("class")
        or request.POST.get("class")
        or ""
    ).strip()

    selected_stream = (
        request.GET.get("stream")
        or request.POST.get("stream")
        or ""
    ).strip()

    selected_window_id = (
        request.GET.get("assessment")
        or request.POST.get("assessment")
        or ""
    ).strip()

    # Streams belonging to selected permitted class.
    streams = sorted(
        {
            stream
            for class_name, stream in allowed_combinations
            if _same_class(
                class_name,
                selected_class,
            )
            and stream
        },
        key=str.casefold,
    )

    # --------------------------------------------------------
    # PERMITTED SUBJECTS FOR SELECTED CLASS/STREAM
    # --------------------------------------------------------

    allowed_subject_ids = _allowed_subject_ids_for_upload(
        request.user,
        selected_class,
        selected_stream,
    )

    subjects = Subject.objects.filter(
        is_active=True,
        id__in=allowed_subject_ids,
    ).order_by("name")

    if request.method == "GET":

        return render(
            request,
            "academic/excel_marks_upload.html",
            {
                "open_windows": open_windows,
                "classes": classes,
                "streams": streams,
                "subjects": subjects,
                "selected_class": selected_class,
                "selected_stream": selected_stream,
                "selected_window_id": selected_window_id,
            },
        )

    # --------------------------------------------------------
    # POST VALUES
    # --------------------------------------------------------

    window_id = request.POST.get("assessment")
    uploaded_file = request.FILES.get("excel_file")

    # --------------------------------------------------------
    # CLASS/STREAM AUTHORIZATION
    # --------------------------------------------------------

    if not selected_class:
        messages.error(
            request,
            "Please select a class.",
        )
        return redirect("academic:excel_upload")

    class_authorized = False

    for class_name, stream in allowed_combinations:

        if not _same_class(
            selected_class,
            class_name,
        ):
            continue

        # A blank assigned stream means whole class.
        if not stream:
            class_authorized = True
            break

        if _same_stream(
            selected_stream,
            stream,
        ):
            class_authorized = True
            break

    if not class_authorized:
        messages.error(
            request,
            "You are not authorized to upload marks for "
            f"{selected_class}"
            f"{' - ' + selected_stream if selected_stream else ''}.",
        )
        return redirect("academic:excel_upload")

    # --------------------------------------------------------
    # REQUIRED FIELDS
    # --------------------------------------------------------

    if not all(
        [
            window_id,
            uploaded_file,
        ]
    ):
        messages.error(
            request,
            "Please complete all fields and select an Excel file.",
        )
        return redirect("academic:excel_upload")

    # --------------------------------------------------------
    # ASSESSMENT WINDOW
    # --------------------------------------------------------

    try:
        window = AssessmentWindow.objects.select_related(
            "academic_year",
            "term",
            "assessment_type",
        ).get(
            pk=window_id,
            is_active=True,
        )
    except AssessmentWindow.DoesNotExist:
        messages.error(
            request,
            "The selected assessment is no longer active in Assessment & Marks Management.",
        )
        return redirect("academic:excel_upload")

    if not window.recording_open():
        messages.error(
            request,
            "This assessment window is closed. Marks cannot be uploaded.",
        )
        return redirect("academic:excel_upload")

    # The selected AssessmentWindow determines the exact
    # academic year, term and assessment type.
    year_id = window.academic_year_id
    term_id = window.term_id
    assessment_type_id = window.assessment_type_id

    # --------------------------------------------------------
    # READ EXCEL
    # --------------------------------------------------------

    try:
        workbook = load_workbook(
            uploaded_file,
            data_only=True,
        )

        sheet = workbook.active

    except Exception as exc:
        messages.error(
            request,
            f"Unable to read the Excel file: {exc}",
        )
        return redirect("academic:excel_upload")

    headers = [
        str(cell.value).strip().casefold()
        if cell.value is not None
        else ""
        for cell in sheet[1]
    ]

    required_headers = [
        "admission no",
        "learner name",
        "subject",
        "mark",
    ]

    if not all(
        header in headers
        for header in required_headers
    ):
        messages.error(
            request,
            "Invalid Excel template. Required columns are: "
            "Admission No, Learner Name, Subject, Mark.",
        )
        return redirect("academic:excel_upload")

    index = {
        header: headers.index(header)
        for header in required_headers
    }

    errors = []
    valid_rows = []
    seen = set()

    # --------------------------------------------------------
    # VALIDATE EVERY ROW
    # --------------------------------------------------------

    for row_number, row in enumerate(
        sheet.iter_rows(
            min_row=2,
            values_only=True,
        ),
        start=2,
    ):

        admission_no = row[index["admission no"]]
        learner_name = row[index["learner name"]]
        excel_subject = row[index["subject"]]
        mark = row[index["mark"]]

        if (
            admission_no is None
            and learner_name is None
            and excel_subject is None
            and mark is None
        ):
            continue

        admission_no = str(
            admission_no
        ).strip()

        try:
            student = Student.objects.get(
                admission_no=admission_no
            )
        except Student.DoesNotExist:
            errors.append(
                f"Row {row_number}: learner "
                f"'{admission_no}' was not found."
            )
            continue

        # ----------------------------------------------------
        # CRITICAL CLASS/STREAM CHECK
        # ----------------------------------------------------

        if not _same_class(
            student.class_name,
            selected_class,
        ):
            errors.append(
                f"Row {row_number}: learner "
                f"{admission_no} belongs to "
                f"{student.class_name}, not "
                f"{selected_class}."
            )
            continue

        if selected_stream and not _same_stream(
            student.stream,
            selected_stream,
        ):
            errors.append(
                f"Row {row_number}: learner "
                f"{admission_no} belongs to stream "
                f"{student.stream or 'Unspecified'}, not "
                f"{selected_stream}."
            )
            continue

        # ----------------------------------------------------
        # SUBJECT CHECK
        # ----------------------------------------------------

        if not excel_subject:
            errors.append(
                f"Row {row_number}: subject is missing."
            )
            continue

        excel_subject_name = (
            str(excel_subject).strip()
            if excel_subject is not None
            else ""
        )

        if not excel_subject_name:
            errors.append(
                f"Row {row_number}: subject is required."
            )
            continue

        try:
            row_subject = Subject.objects.get(
                name__iexact=excel_subject_name,
                is_active=True,
            )
        except Subject.DoesNotExist:
            errors.append(
                f"Row {row_number}: subject "
                f"'{excel_subject_name}' was not found."
            )
            continue
        except Subject.MultipleObjectsReturned:
            errors.append(
                f"Row {row_number}: subject "
                f"'{excel_subject_name}' is ambiguous."
            )
            continue

        if row_subject.id not in allowed_subject_ids:
            errors.append(
                f"Row {row_number}: you are not authorized "
                f"to enter {row_subject.name} marks for "
                f"{student.class_name}"
                f"{' - ' + student.stream if student.stream else ''}."
            )
            continue

        # ----------------------------------------------------
        # MARK CHECK
        # ----------------------------------------------------

        if mark is None or str(mark).strip() == "":
            errors.append(
                f"Row {row_number}: mark is missing."
            )
            continue

        try:
            numeric_mark = float(mark)
        except (TypeError, ValueError):
            errors.append(
                f"Row {row_number}: '{mark}' is not a valid mark."
            )
            continue

        if numeric_mark < 0 or numeric_mark > 100:
            errors.append(
                f"Row {row_number}: mark must be between 0 and 100."
            )
            continue

        # ----------------------------------------------------
        # DUPLICATE CHECK
        # ----------------------------------------------------

        key = (
            student.pk,
            row_subject.pk,
        )

        if key in seen:
            errors.append(
                f"Row {row_number}: duplicate mark for "
                f"learner '{admission_no}' and subject "
                f"'{row_subject.name}'."
            )
            continue

        seen.add(key)

        # ----------------------------------------------------
        # FINAL SERVER-SIDE PERMISSION CHECK
        # ----------------------------------------------------

        if not can_enter_marks(
            request.user,
            student.class_name,
            student.stream,
            row_subject,
        ):
            errors.append(
                f"Row {row_number}: you are not authorized "
                f"to enter {row_subject.name} marks for "
                f"{student.class_name}"
                f"{' - ' + student.stream if student.stream else ''}."
            )
            continue

        valid_rows.append(
            {
                "student": student,
                "subject": row_subject,
                "score": numeric_mark,
            }
        )

    # --------------------------------------------------------
    # NEVER SAVE PARTIAL FILE
    # --------------------------------------------------------

    if errors:
        return render(
            request,
            "academic/excel_marks_upload.html",
            {
                "open_windows": open_windows,
                "years": years,
                "terms": terms,
                "assessment_types": assessment_types,
                "classes": classes,
                "streams": streams,
                "subjects": subjects,
                "errors": errors,
                "valid_rows": valid_rows,
                "selected_year": year_id,
                "selected_term": term_id,
                "selected_assessment_type": assessment_type_id,
                "selected_class": selected_class,
                "selected_stream": selected_stream,
            },
        )

    if not valid_rows:
        messages.warning(
            request,
            "No valid learner marks were found in the uploaded file.",
        )
        return redirect("academic:excel_upload")

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    for item in valid_rows:

        Assessment.objects.update_or_create(
            student=item["student"],
            academic_year_id=year_id,
            term_id=term_id,
            subject=item["subject"],
            assessment_type_id=assessment_type_id,
            defaults={
                "score": item["score"],
            },
        )

    messages.success(
        request,
        f"{len(valid_rows)} mark(s) uploaded successfully.",
    )

    return redirect("academic:marks_entry")



@login_required
def upload_individual_marks_excel(request):
    """
    Upload an Individual Learner Excel workbook.

    Expected format:
        Admission No | Learner Name | Subject | Mark

    The learner and assessment are taken from the workbook rows and
    the selected AssessmentWindow. Teacher permissions are enforced
    for every subject before anything is saved.
    """

    open_windows = [
        window
        for window in AssessmentWindow.objects.filter(
            is_active=True
        ).select_related(
            "academic_year",
            "term",
            "assessment_type",
        )
        if window.recording_open()
    ]

    selected_window_id = (
        request.GET.get("assessment")
        or request.POST.get("assessment")
        or ""
    ).strip()

    if request.method == "GET":

        return render(
            request,
            "academic/excel_individual_marks_upload.html",
            {
                "open_windows": open_windows,
                "selected_window_id": selected_window_id,
            },
        )

    # --------------------------------------------------------
    # REQUIRED VALUES
    # --------------------------------------------------------

    window_id = request.POST.get("assessment")
    uploaded_file = request.FILES.get("excel_file")

    if not window_id or not uploaded_file:
        messages.error(
            request,
            "Please select an assessment and Excel file.",
        )
        return redirect("academic:excel_upload_individual")

    # --------------------------------------------------------
    # ASSESSMENT WINDOW
    # --------------------------------------------------------

    try:
        window = AssessmentWindow.objects.select_related(
            "academic_year",
            "term",
            "assessment_type",
        ).get(
            pk=window_id,
            is_active=True,
        )
    except AssessmentWindow.DoesNotExist:
        messages.error(
            request,
            "The selected assessment is no longer active.",
        )
        return redirect("academic:excel_upload_individual")

    if not window.recording_open():
        messages.error(
            request,
            "This assessment window is closed. Marks cannot be uploaded.",
        )
        return redirect("academic:excel_upload_individual")

    # --------------------------------------------------------
    # READ WORKBOOK
    # --------------------------------------------------------

    try:
        workbook = load_workbook(
            uploaded_file,
            data_only=True,
        )
        sheet = workbook["Marks Upload"]
    except KeyError:
        messages.error(
            request,
            "Invalid Individual Learner Excel file. "
            "The workbook must contain a 'Marks Upload' worksheet.",
        )
        return redirect("academic:excel_upload_individual")
    except Exception as exc:
        messages.error(
            request,
            f"Unable to read the Excel file: {exc}",
        )
        return redirect("academic:excel_upload_individual")

    # --------------------------------------------------------
    # VERIFY TEMPLATE
    # --------------------------------------------------------

    expected_headers = [
        "Admission No",
        "Learner Name",
        "Subject",
        "Mark",
    ]

    actual_headers = [
        str(cell.value).strip()
        if cell.value is not None
        else ""
        for cell in sheet[1]
    ]

    if actual_headers[:4] != expected_headers:
        messages.error(
            request,
            "Invalid Individual Learner Excel template. "
            "Required columns are: "
            "Admission No, Learner Name, Subject, Mark.",
        )
        return redirect("academic:excel_upload_individual")

    # --------------------------------------------------------
    # PROCESS ALL ROWS FIRST
    # NOTHING IS SAVED UNTIL EVERYTHING IS VALID.
    # --------------------------------------------------------

    pending_marks = []
    errors = []
    seen = set()
    learner_ids = set()

    active_subjects = list(
        Subject.objects.filter(
            is_active=True
        ).order_by("name")
    )

    subjects_by_name = {
        str(subject.name).strip().casefold(): subject
        for subject in active_subjects
    }

    for row_number, row in enumerate(
        sheet.iter_rows(
            min_row=2,
            values_only=True,
        ),
        start=2,
    ):

        if not row:
            continue

        admission_no = (
            str(row[0]).strip()
            if len(row) > 0 and row[0] is not None
            else ""
        )

        learner_name = (
            str(row[1]).strip()
            if len(row) > 1 and row[1] is not None
            else ""
        )

        subject_name = (
            str(row[2]).strip()
            if len(row) > 2 and row[2] is not None
            else ""
        )

        mark_value = (
            row[3]
            if len(row) > 3
            else None
        )

        # Completely blank row.
        if not admission_no and not learner_name and not subject_name and mark_value is None:
            continue

        if not admission_no:
            errors.append(
                f"Row {row_number}: Admission No is required."
            )
            continue

        if not subject_name:
            errors.append(
                f"Row {row_number}: Subject is required."
            )
            continue

        # Blank marks are allowed and are not uploaded.
        if mark_value is None or str(mark_value).strip() == "":
            continue

        # ----------------------------------------------------
        # LEARNER
        # ----------------------------------------------------

        try:
            student = Student.objects.get(
                admission_no__iexact=admission_no
            )
        except Student.DoesNotExist:
            errors.append(
                f"Row {row_number}: Learner with Admission No "
                f"'{admission_no}' was not found."
            )
            continue
        except Student.MultipleObjectsReturned:
            errors.append(
                f"Row {row_number}: Admission No "
                f"'{admission_no}' is not unique."
            )
            continue

        learner_ids.add(student.pk)

        # ----------------------------------------------------
        # CLASS ACCESS
        # ----------------------------------------------------

        student_stream = getattr(
            student,
            "stream",
            None,
        )

        if not can_enter_marks(
            request.user,
            student.class_name,
            student_stream,
            None,
        ):
            # Some permission implementations require a subject,
            # so class access is checked separately where possible.
            from .mark_permissions import can_access_class

            if not can_access_class(
                request.user,
                student.class_name,
                student_stream,
            ):
                errors.append(
                    f"Row {row_number}: You are not authorized "
                    f"to enter marks for learner '{admission_no}'."
                )
                continue

        # ----------------------------------------------------
        # SUBJECT
        # ----------------------------------------------------

        subject = subjects_by_name.get(
            subject_name.casefold()
        )

        if not subject:
            errors.append(
                f"Row {row_number}: Subject '{subject_name}' "
                f"is not an active subject."
            )
            continue

        # ----------------------------------------------------
        # SUBJECT PERMISSION
        # ----------------------------------------------------

        if not can_enter_marks(
            request.user,
            student.class_name,
            student_stream,
            subject,
        ):
            errors.append(
                f"Row {row_number}: You are not authorized "
                f"to enter {subject.name} for "
                f"{admission_no}."
            )
            continue

        # ----------------------------------------------------
        # MARK VALIDATION
        # ----------------------------------------------------

        try:
            numeric_mark = float(mark_value)
        except (TypeError, ValueError):
            errors.append(
                f"Row {row_number}: Mark for {subject.name} "
                f"must be a number between 0 and 100."
            )
            continue

        if numeric_mark < 0 or numeric_mark > 100:
            errors.append(
                f"Row {row_number}: Mark for {subject.name} "
                f"must be between 0 and 100."
            )
            continue

        # Convert whole-number Excel values cleanly.
        if numeric_mark.is_integer():
            numeric_mark = int(numeric_mark)

        duplicate_key = (
            student.pk,
            subject.pk,
        )

        if duplicate_key in seen:
            errors.append(
                f"Row {row_number}: Duplicate mark for "
                f"{admission_no} - {subject.name}."
            )
            continue

        seen.add(duplicate_key)

        pending_marks.append(
            (
                student,
                subject,
                numeric_mark,
            )
        )

    # --------------------------------------------------------
    # VERIFY THIS IS AN INDIVIDUAL LEARNER FILE
    # --------------------------------------------------------

    if len(learner_ids) > 1:
        errors.append(
            "This Individual Learner Excel file contains "
            "more than one learner. Please upload a file "
            "generated for one learner only."
        )

    if not pending_marks and not errors:
        errors.append(
            "No marks were found in the Excel file."
        )

    # --------------------------------------------------------
    # ALL OR NOTHING
    # --------------------------------------------------------

    if errors:
        for error in errors[:15]:
            messages.error(request, error)

        if len(errors) > 15:
            messages.error(
                request,
                f"{len(errors) - 15} additional validation errors "
                "were found.",
            )

        return redirect("academic:excel_upload_individual")

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    saved = 0

    for student, subject, score in pending_marks:

        Assessment.objects.update_or_create(
            student=student,
            academic_year_id=window.academic_year_id,
            term_id=window.term_id,
            subject=subject,
            assessment_type_id=window.assessment_type_id,
            defaults={
                "score": score,
            },
        )

        saved += 1

    messages.success(
        request,
        f"Individual learner Excel uploaded successfully. "
        f"{saved} mark(s) saved.",
    )

    return redirect("academic:marks_entry")

@login_required
def upload_class_marks_excel(request):
    """
    Upload the Whole Class matrix Excel format.

    Format:
        Admission No | Learner Name | Stream | Subject 1 | Subject 2 | ...

    Learners are rows.
    Subjects are columns.
    """

    open_windows = [
        window
        for window in AssessmentWindow.objects.filter(
            is_active=True
        ).select_related(
            "academic_year",
            "term",
            "assessment_type",
        )
        if window.recording_open()
    ]

    allowed_combinations = _allowed_class_streams_for_upload(
        request.user
    )

    classes = sorted(
        {
            class_name
            for class_name, stream in allowed_combinations
        },
        key=str.casefold,
    )

    selected_class = (
        request.GET.get("class")
        or request.POST.get("class")
        or ""
    ).strip()

    selected_stream = (
        request.GET.get("stream")
        or request.POST.get("stream")
        or ""
    ).strip()

    selected_window_id = (
        request.GET.get("assessment")
        or request.POST.get("assessment")
        or ""
    ).strip()

    streams = sorted(
        {
            stream
            for class_name, stream in allowed_combinations
            if _same_class(
                class_name,
                selected_class,
            )
            and stream
        },
        key=str.casefold,
    )

    if request.method == "GET":

        return render(
            request,
            "academic/excel_marks_upload.html",
            {
                "open_windows": open_windows,
                "classes": classes,
                "streams": streams,
                "subjects": Subject.objects.none(),
                "selected_class": selected_class,
                "selected_stream": selected_stream,
                "selected_window_id": selected_window_id,
                "whole_class_matrix": True,
            },
        )

    # --------------------------------------------------------
    # REQUIRED VALUES
    # --------------------------------------------------------

    window_id = request.POST.get("assessment")
    uploaded_file = request.FILES.get("excel_file")

    if not selected_class:
        messages.error(request, "Please select a class.")
        return redirect("academic:excel_upload_class")

    if not window_id or not uploaded_file:
        messages.error(
            request,
            "Please select an assessment and Excel file.",
        )
        return redirect("academic:excel_upload_class")

    # --------------------------------------------------------
    # CLASS AUTHORIZATION
    # --------------------------------------------------------

    class_authorized = False

    for class_name, stream in allowed_combinations:

        if not _same_class(
            selected_class,
            class_name,
        ):
            continue

        if not stream:
            class_authorized = True
            break

        if _same_stream(
            selected_stream,
            stream,
        ):
            class_authorized = True
            break

    if not class_authorized:
        messages.error(
            request,
            "You are not authorized to upload marks for "
            f"{selected_class}"
            f"{' - ' + selected_stream if selected_stream else ''}.",
        )
        return redirect("academic:excel_upload_class")

    # --------------------------------------------------------
    # ASSESSMENT WINDOW
    # --------------------------------------------------------

    try:
        window = AssessmentWindow.objects.select_related(
            "academic_year",
            "term",
            "assessment_type",
        ).get(
            pk=window_id,
            is_active=True,
        )
    except AssessmentWindow.DoesNotExist:
        messages.error(
            request,
            "The selected assessment is no longer active.",
        )
        return redirect("academic:excel_upload_class")

    if not window.recording_open():
        messages.error(
            request,
            "This assessment window is closed. Marks cannot be uploaded.",
        )
        return redirect("academic:excel_upload_class")

    # --------------------------------------------------------
    # READ WORKBOOK
    # --------------------------------------------------------

    try:
        workbook = load_workbook(
            uploaded_file,
            data_only=True,
        )
        sheet = workbook.active
    except Exception as exc:
        messages.error(
            request,
            f"Unable to read the Excel file: {exc}",
        )
        return redirect("academic:excel_upload_class")

    headers = [
        str(cell.value).strip()
        if cell.value is not None
        else ""
        for cell in sheet[1]
    ]

    normalized = [
        value.casefold()
        for value in headers
    ]

    required_identity_headers = [
        "admission no",
        "learner name",
        "stream",
    ]

    missing_identity = [
        h
        for h in required_identity_headers
        if h not in normalized
    ]

    if missing_identity:
        messages.error(
            request,
            "Invalid Whole Class Excel format. Missing: "
            + ", ".join(missing_identity),
        )
        return redirect("academic:excel_upload_class")

    admission_index = normalized.index("admission no")
    name_index = normalized.index("learner name")
    stream_index = normalized.index("stream")

    # Every column after Stream is a subject.
    subject_columns = []

    for index in range(3, len(headers)):

        header = headers[index].strip()

        if not header:
            continue

        try:
            row_subject = Subject.objects.get(
                name__iexact=header,
                is_active=True,
            )
        except Subject.DoesNotExist:
            messages.error(
                request,
                f"Invalid subject column '{header}'. "
                "The subject does not exist or is inactive.",
            )
            return redirect("academic:excel_upload_class")
        except Subject.MultipleObjectsReturned:
            messages.error(
                request,
                f"Subject column '{header}' is ambiguous.",
            )
            return redirect("academic:excel_upload_class")

        subject_columns.append(
            (index, row_subject)
        )

    if not subject_columns:
        messages.error(
            request,
            "No subject columns were found after Stream.",
        )
        return redirect("academic:excel_upload_class")

    # --------------------------------------------------------
    # VALIDATE ALL ROWS BEFORE SAVING ANYTHING
    # --------------------------------------------------------

    errors = []
    valid_rows = []
    seen = set()

    for row_number, row in enumerate(
        sheet.iter_rows(
            min_row=2,
            values_only=True,
        ),
        start=2,
    ):

        admission_no = row[admission_index]

        if admission_no is None or str(admission_no).strip() == "":
            continue

        admission_no = str(admission_no).strip()

        try:
            student = Student.objects.get(
                admission_no=admission_no
            )
        except Student.DoesNotExist:
            errors.append(
                f"Row {row_number}: learner "
                f"'{admission_no}' was not found."
            )
            continue

        if not _same_class(
            student.class_name,
            selected_class,
        ):
            errors.append(
                f"Row {row_number}: learner {admission_no} "
                f"belongs to {student.class_name}, not "
                f"{selected_class}."
            )
            continue

        actual_stream = (
            str(getattr(student, "stream", "") or "").strip()
        )

        excel_stream = (
            str(row[stream_index] or "").strip()
        )

        if selected_stream:
            if not _same_stream(
                actual_stream,
                selected_stream,
            ):
                errors.append(
                    f"Row {row_number}: learner {admission_no} "
                    f"is not in stream {selected_stream}."
                )
                continue

        if excel_stream and not _same_stream(
            actual_stream,
            excel_stream,
        ):
            errors.append(
                f"Row {row_number}: learner {admission_no} "
                f"has stream {actual_stream or 'Unspecified'}, "
                f"not {excel_stream}."
            )
            continue

        for column_index, row_subject in subject_columns:

            raw_mark = row[column_index]

            # Blank means no mark entered.
            if raw_mark is None or str(raw_mark).strip() == "":
                continue

            if not can_enter_marks(
                request.user,
                student.class_name,
                actual_stream or None,
                row_subject,
            ):
                errors.append(
                    f"Row {row_number}: you are not authorized "
                    f"to enter {row_subject.name} marks for "
                    f"{student.class_name}"
                    f"{' - ' + actual_stream if actual_stream else ''}."
                )
                continue

            try:
                numeric_mark = float(raw_mark)
            except (TypeError, ValueError):
                errors.append(
                    f"Row {row_number}: {row_subject.name} "
                    f"contains an invalid mark."
                )
                continue

            if numeric_mark < 0 or numeric_mark > 100:
                errors.append(
                    f"Row {row_number}: {row_subject.name} "
                    "mark must be between 0 and 100."
                )
                continue

            key = (
                student.pk,
                row_subject.pk,
            )

            if key in seen:
                errors.append(
                    f"Row {row_number}: duplicate "
                    f"{row_subject.name} mark for "
                    f"{admission_no}."
                )
                continue

            seen.add(key)

            valid_rows.append(
                {
                    "student": student,
                    "subject": row_subject,
                    "score": numeric_mark,
                }
            )

    # --------------------------------------------------------
    # NEVER SAVE PARTIAL FILE
    # --------------------------------------------------------

    if errors:
        messages.error(
            request,
            "Upload rejected. No marks were saved.",
        )

        for error in errors[:15]:
            messages.error(
                request,
                error,
            )

        return redirect(
            "academic:excel_upload_class"
        )

    if not valid_rows:
        messages.warning(
            request,
            "No marks were entered in the uploaded workbook.",
        )
        return redirect(
            "academic:excel_upload_class"
        )

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    for item in valid_rows:

        Assessment.objects.update_or_create(
            student=item["student"],
            academic_year_id=window.academic_year_id,
            term_id=window.term_id,
            subject=item["subject"],
            assessment_type_id=window.assessment_type_id,
            defaults={
                "score": item["score"],
            },
        )

    messages.success(
        request,
        f"{len(valid_rows)} Whole Class mark(s) uploaded successfully.",
    )

    return redirect("academic:marks_entry")


